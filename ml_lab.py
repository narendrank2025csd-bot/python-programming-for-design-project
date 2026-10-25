import hashlib
import json
import math
import os
import re
import uuid
from zipfile import BadZipFile
from functools import lru_cache
from pathlib import Path

import joblib
import numpy as np
import pandas as pd
from openpyxl import load_workbook as load_xlsx_read_only
from sklearn.cluster import KMeans
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LinearRegression, LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    confusion_matrix,
    f1_score,
    mean_absolute_error,
    mean_squared_error,
    precision_score,
    r2_score,
    recall_score,
    silhouette_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler
from sklearn.tree import DecisionTreeClassifier, DecisionTreeRegressor
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor


BASE_DIR = Path(__file__).resolve().parent
WORKBOOK_PATH = BASE_DIR / "data" / "Startup_Tycoon_Combined_ML_Dataset_Pack.xlsx"
ARTIFACT_DIR = BASE_DIR / "instance" / "ml_models"
MAX_DATA_ROWS = 20000
MAX_PREVIEW_ROWS = 50
MAX_TRAIN_FEATURES = 20
DATASET_META = {
    "Synthetic_Sales_Demo": {
        "description": "Generated example sales records for demonstrating regression and exploratory analysis.",
        "origin": "Synthetic records included in the supplied workbook.",
        "synthetic": True,
        "tasks": ["regression", "classification", "clustering"],
    },
    "Synthetic_India_Prices": {
        "description": "Generated daily commodity price examples for forecasting prototypes and trend exploration.",
        "origin": "Synthetic illustrative records included in the supplied workbook; not observed Indian market prices.",
        "synthetic": True,
        "tasks": ["regression", "classification", "clustering"],
    },
    "Synthetic_Ecommerce_Survey": {
        "description": "Generated survey-like examples for classification and exploratory analysis.",
        "origin": "Synthetic survey-like records included in the supplied workbook; not responses from real people.",
        "synthetic": True,
        "tasks": ["regression", "classification", "clustering"],
    },
}
ALGORITHMS = {
    "regression": {
        "linear_regression": LinearRegression,
        "decision_tree_regressor": DecisionTreeRegressor,
        "random_forest_regressor": RandomForestRegressor,
    },
    "classification": {
        "logistic_regression": LogisticRegression,
        "decision_tree_classifier": DecisionTreeClassifier,
        "random_forest_classifier": RandomForestClassifier,
    },
    "clustering": {"kmeans": KMeans},
}
ALGORITHM_LABELS = {
    "linear_regression": "Linear Regression",
    "decision_tree_regressor": "Decision Tree Regressor",
    "random_forest_regressor": "Random Forest Regressor",
    "logistic_regression": "Logistic Regression",
    "decision_tree_classifier": "Decision Tree Classifier",
    "random_forest_classifier": "Random Forest Classifier",
    "kmeans": "K-Means",
}
ML_QUESTS = [
    {"id": "ml-explore", "title": "Explore Your First Dataset", "description": "Open a workbook dataset in the Research Lab.", "event": "explore", "difficulty": "Easy"},
    {"id": "ml-analysis", "title": "Complete a Dataset Analysis", "description": "Inspect real column statistics or generate a dataset chart.", "event": "analysis", "difficulty": "Easy"},
    {"id": "ml-regression", "title": "Train Your First Regression Model", "description": "Successfully train a regression pipeline and evaluate its held-out predictions.", "event": "regression", "difficulty": "Medium"},
    {"id": "ml-classification", "title": "Train a Classification Model", "description": "Successfully train a classification pipeline and evaluate its held-out predictions.", "event": "classification", "difficulty": "Medium"},
    {"id": "ml-evaluation", "title": "Evaluate a Model", "description": "Complete an experiment with real evaluation metrics.", "event": "evaluation", "difficulty": "Easy"},
    {"id": "ml-compare", "title": "Compare Two Models", "description": "Compare actual held-out evaluation metrics from two compatible experiments.", "event": "compare", "difficulty": "Medium"},
    {"id": "ml-save", "title": "Save Your First Model", "description": "Save a trained model to your personal model collection.", "event": "save", "difficulty": "Easy"},
    {"id": "ml-predict", "title": "Generate a Real Prediction", "description": "Run a prediction with a model you own.", "event": "prediction", "difficulty": "Easy"},
]


class MLLabError(ValueError):
    pass


def _workbook_fingerprint():
    if not WORKBOOK_PATH.is_file():
        raise MLLabError(
            "The supplied workbook is missing. Place it at "
            "data/Startup_Tycoon_Combined_ML_Dataset_Pack.xlsx."
        )
    digest = hashlib.sha256()
    with WORKBOOK_PATH.open("rb") as workbook:
        for chunk in iter(lambda: workbook.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@lru_cache(maxsize=2)
def _load_workbook(path, modified_ns, size):
    del modified_ns
    if size > 25 * 1024 * 1024:
        raise MLLabError("The workbook exceeds the 25 MB processing limit.")
    try:
        expected = {
            "READ_ME_FIRST",
            "Synthetic_Sales_Demo",
            "Synthetic_India_Prices",
            "Synthetic_Ecommerce_Survey",
            "Source_Catalog",
        }
        workbook = load_xlsx_read_only(path, read_only=True, data_only=True)
        try:
            if len(workbook.sheetnames) > 20 or not expected.issubset(set(workbook.sheetnames)):
                raise MLLabError("The workbook is missing required sheets or exceeds the sheet limit.")
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) > MAX_DATA_ROWS + 1 or (sheet.max_column or 0) > 200:
                    raise MLLabError(f"{sheet.title} exceeds the safe workbook row or column limit.")
        finally:
            workbook.close()
        frames = {}
        with pd.ExcelFile(path, engine="openpyxl") as excel:
            if len(excel.sheet_names) > 20 or not expected.issubset(set(excel.sheet_names)):
                raise MLLabError("The workbook is missing required sheets or exceeds the sheet limit.")
            for sheet in excel.sheet_names:
                frame = pd.read_excel(excel, sheet_name=sheet)
                if len(frame) > MAX_DATA_ROWS and sheet in DATASET_META:
                    raise MLLabError(f"{sheet} exceeds the {MAX_DATA_ROWS:,}-row training limit.")
                if frame.shape[1] > 200:
                    raise MLLabError(f"{sheet} exceeds the safe column limit.")
                frame.columns = [str(column).strip() for column in frame.columns]
                if any(not column for column in frame.columns) or len(set(frame.columns)) != len(frame.columns):
                    raise MLLabError(f"{sheet} has blank or duplicate column names.")
                frames[sheet] = frame
    except MLLabError:
        raise
    except (OSError, ValueError, BadZipFile, ImportError) as error:
        raise MLLabError("The supplied file could not be opened as a valid Excel workbook.") from error
    return frames


def load_workbook():
    stat = WORKBOOK_PATH.stat() if WORKBOOK_PATH.is_file() else None
    return _load_workbook(
        str(WORKBOOK_PATH),
        stat.st_mtime_ns if stat else 0,
        stat.st_size if stat else 0,
    )


def dataset_fingerprint():
    return _workbook_fingerprint()


def _json_value(value):
    if value is None:
        return None
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, (np.bool_,)):
        return bool(value)
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _column_type(series):
    if pd.api.types.is_datetime64_any_dtype(series):
        return "date"
    if pd.api.types.is_numeric_dtype(series):
        return "numeric"
    return "categorical"


def _column_summaries(frame):
    summaries = []
    numeric_frame = frame.select_dtypes(include=[np.number])
    numeric_summary = (
        numeric_frame.describe().transpose().to_dict("index")
        if not numeric_frame.empty
        else {}
    )
    for column in frame.columns:
        values = frame[column]
        counts = values.value_counts(dropna=True)
        stats = numeric_summary.get(column, {})
        summaries.append(
            {
                "name": column,
                "dtype": str(values.dtype),
                "kind": _column_type(values),
                "missing": int(values.isna().sum()),
                "unique": int(values.nunique(dropna=True)),
                "samples": [_json_value(item) for item in counts.index[:12]],
                "top_values": [
                    {"label": str(item), "count": int(count)}
                    for item, count in counts.head(10).items()
                ],
                "statistics": {
                    key: _json_value(stats.get(key))
                    for key in ("mean", "std", "min", "25%", "50%", "75%", "max")
                },
            }
        )
    return summaries


def dataset_catalog():
    frames = load_workbook()
    fingerprint = _workbook_fingerprint()
    entries = []
    for sheet in frames:
        if sheet in DATASET_META:
            meta = DATASET_META[sheet]
            frame = frames[sheet]
            entries.append(
                {
                    "name": sheet,
                    "description": meta["description"],
                    "rows": len(frame),
                    "features": len(frame.columns),
                    "columns": _column_summaries(frame),
                    "origin": meta["origin"],
                    "synthetic": meta["synthetic"],
                    "tasks": meta["tasks"],
                    "trainable": True,
                    "version": fingerprint[:16],
                }
            )
        elif sheet == "Source_Catalog":
            entries.append(
                {
                    "name": sheet,
                    "description": "References to external real datasets. Those source datasets are not included or imported here.",
                    "rows": len(frames[sheet]),
                    "features": len(frames[sheet].columns),
                    "columns": _column_summaries(frames[sheet]),
                    "origin": "Reference catalog supplied in the workbook; not imported source observations.",
                    "synthetic": False,
                    "tasks": [],
                    "trainable": False,
                    "version": fingerprint[:16],
                }
            )
        else:
            frame = frames[sheet]
            entries.append(
                {
                    "name": sheet,
                    "description": "Workbook guidance and usage notes.",
                    "rows": len(frame),
                    "features": len(frame.columns),
                    "columns": _column_summaries(frame),
                    "origin": "Workbook documentation.",
                    "synthetic": False,
                    "tasks": [],
                    "trainable": False,
                    "version": fingerprint[:16],
                }
            )
    return entries


def explore_dataset(name, preview_rows=20):
    if name not in DATASET_META and name not in {"READ_ME_FIRST", "Source_Catalog"}:
        raise MLLabError("Select a dataset from the supplied workbook.")
    frames = load_workbook()
    frame = frames[name]
    summaries = _column_summaries(frame)
    numeric = frame.select_dtypes(include=[np.number])
    categorical = [column for column in frame.columns if column not in numeric.columns]
    return {
        "name": name,
        "rows": len(frame),
        "features": len(frame.columns),
        "columns": summaries,
        "preview_columns": [str(column) for column in frame.columns],
        "preview": [
            {str(key): _json_value(value) for key, value in row.items()}
            for row in frame.head(min(max(int(preview_rows), 1), MAX_PREVIEW_ROWS)).to_dict("records")
        ],
        "numeric_columns": list(numeric.columns),
        "categorical_columns": categorical,
        "synthetic": name in DATASET_META and DATASET_META[name]["synthetic"],
        "trainable": name in DATASET_META,
        "fingerprint": _workbook_fingerprint(),
    }


def chart_data(dataset_name, column_name):
    if dataset_name not in DATASET_META and dataset_name not in {"READ_ME_FIRST", "Source_Catalog"}:
        raise MLLabError("Select a dataset from the supplied workbook.")
    frame = load_workbook()[dataset_name]
    if column_name not in frame.columns:
        raise MLLabError("Select a valid column from this dataset.")
    column = frame[column_name]
    if pd.api.types.is_numeric_dtype(column):
        clean = column.dropna()
        if clean.empty:
            raise MLLabError("This column has no numeric values to chart.")
        bins = min(20, max(5, int(np.sqrt(len(clean)))))
        counts, edges = np.histogram(clean.to_numpy(dtype=float), bins=bins)
        labels = [f"{edges[i]:.3g}–{edges[i + 1]:.3g}" for i in range(len(counts))]
        return {"kind": "numeric", "labels": labels, "values": counts.astype(int).tolist()}
    counts = column.fillna("(missing)").astype(str).value_counts().head(15)
    return {"kind": "categorical", "labels": counts.index.tolist(), "values": counts.astype(int).tolist()}


def _prepare_dates(frame):
    converted = frame.copy()
    for column in converted.columns:
        if column.lower() == "date" or column.lower().endswith("_date"):
            parsed = pd.to_datetime(converted[column], errors="coerce", format="mixed")
            if parsed.notna().sum() == converted[column].notna().sum() and parsed.notna().any():
                converted[column] = (parsed.astype("int64") // 86_400_000_000_000).astype(float)
                converted.loc[parsed.isna(), column] = np.nan
    for column in converted.select_dtypes(include=[np.number]).columns:
        converted[column] = converted[column].replace([np.inf, -np.inf], np.nan)
    return converted


def _preprocessor(features):
    numeric = [name for name in features.columns if pd.api.types.is_numeric_dtype(features[name])]
    categorical = [name for name in features.columns if name not in numeric]
    transformers = []
    if numeric:
        transformers.append(
            (
                "numeric",
                Pipeline(
                    [
                        ("imputer", SimpleImputer(strategy="median")),
                        ("scaler", StandardScaler()),
                    ]
                ),
                numeric,
            )
        )
    if categorical:
        transformers.append(
            (
                "categorical",
                Pipeline(
                    [
                        ("imputer", SimpleImputer(strategy="most_frequent")),
                        ("encoder", OneHotEncoder(handle_unknown="ignore", min_frequency=2)),
                    ]
                ),
                categorical,
            )
        )
    if not transformers:
        raise MLLabError("Choose at least one usable input feature.")
    return ColumnTransformer(transformers, remainder="drop")


def _model(task, algorithm, parameters):
    if task not in ALGORITHMS or algorithm not in ALGORITHMS[task]:
        raise MLLabError("That algorithm is not compatible with the selected task.")
    seed = 42
    if algorithm == "linear_regression":
        return LinearRegression()
    if algorithm == "logistic_regression":
        return LogisticRegression(max_iter=1000, random_state=seed)
    if algorithm in {"decision_tree_regressor", "decision_tree_classifier"}:
        depth = parameters.get("max_depth")
        max_depth = int(depth) if depth not in (None, "") else None
        if max_depth is not None and not 1 <= max_depth <= 40:
            raise MLLabError("Tree depth must be between 1 and 40.")
        estimator = DecisionTreeRegressor if task == "regression" else DecisionTreeClassifier
        return estimator(max_depth=max_depth, random_state=seed)
    if algorithm in {"random_forest_regressor", "random_forest_classifier"}:
        estimators = int(parameters.get("n_estimators", 100))
        if not 10 <= estimators <= 250:
            raise MLLabError("Forest size must be between 10 and 250 trees.")
        depth = parameters.get("max_depth")
        max_depth = int(depth) if depth not in (None, "") else None
        if max_depth is not None and not 1 <= max_depth <= 40:
            raise MLLabError("Tree depth must be between 1 and 40.")
        estimator = RandomForestRegressor if task == "regression" else RandomForestClassifier
        return estimator(n_estimators=estimators, max_depth=max_depth, random_state=seed, n_jobs=1)
    clusters = int(parameters.get("n_clusters", 3))
    if not 2 <= clusters <= 10:
        raise MLLabError("K-Means must use between 2 and 10 clusters.")
    return KMeans(n_clusters=clusters, random_state=seed, n_init=10)


def train_model(dataset_name, task, target, feature_names, algorithm, parameters, test_size=0.2):
    if dataset_name not in DATASET_META:
        raise MLLabError("Only the three synthetic example sheets contain model-training data.")
    frame = load_workbook()[dataset_name]
    if len(frame) > MAX_DATA_ROWS:
        raise MLLabError(f"Training is limited to {MAX_DATA_ROWS:,} rows.")
    if task not in ALGORITHMS:
        raise MLLabError("Select regression, classification, or clustering.")
    try:
        test_size = float(test_size)
    except (TypeError, ValueError):
        raise MLLabError("Test set size must be a number between 0.1 and 0.4.")
    if not 0.1 <= test_size <= 0.4:
        raise MLLabError("Test set size must be between 10% and 40%.")
    if not isinstance(feature_names, list):
        raise MLLabError("Choose input features explicitly.")
    feature_names = list(dict.fromkeys(feature_names))
    if not feature_names or len(feature_names) > MAX_TRAIN_FEATURES:
        raise MLLabError(f"Choose between 1 and {MAX_TRAIN_FEATURES} input features.")
    if any(not isinstance(name, str) or name not in frame.columns for name in feature_names):
        raise MLLabError("One or more selected features do not exist in this dataset.")
    if task == "clustering":
        target = ""
    else:
        if not isinstance(target, str) or target not in frame.columns:
            raise MLLabError("Select a valid target column.")
        if target in feature_names:
            raise MLLabError("The target cannot also be an input feature.")
    identifier_features = [
        name for name in feature_names
        if re.search(r"(^|[_\s])id($|[_\s])", name, flags=re.IGNORECASE)
        or name.lower().endswith("identifier")
    ]
    warnings = []
    if identifier_features:
        warnings.append(
            "Identifier-like columns were explicitly selected: "
            + ", ".join(identifier_features)
            + ". They may cause leakage or meaningless patterns."
        )
    if dataset_name == "Synthetic_Sales_Demo" and target == "sales_revenue":
        derived_inputs = {"price", "units_sold"}.intersection(feature_names)
        if derived_inputs:
            warnings.append(
                "sales_revenue is constructed from price × units_sold in this synthetic sheet; "
                "selected inputs " + ", ".join(sorted(derived_inputs))
                + " may reveal the target formula rather than generalize."
            )
    model = _model(task, algorithm, parameters or {})
    if task == "clustering" and len(feature_names) > 12:
        raise MLLabError("K-Means is limited to 12 selected features.")
    selected = _prepare_dates(frame[feature_names])
    if selected.isna().all(axis=None):
        raise MLLabError("The selected input columns contain no usable values.")
    for name in selected.columns:
        if selected[name].isna().all():
            raise MLLabError(f"Feature {name} contains no usable values.")
    if task == "clustering":
        if len(frame) < 10:
            raise MLLabError("K-Means needs at least 10 rows for a useful exploratory result.")
        pipeline = Pipeline([("preprocessor", _preprocessor(selected)), ("model", model)])
        assignments = pipeline.fit_predict(selected)
        sizes = np.bincount(assignments, minlength=model.n_clusters)
        silhouette = None
        if 1 < len(np.unique(assignments)) < len(assignments):
            try:
                silhouette = float(
                    silhouette_score(
                        pipeline.named_steps["preprocessor"].transform(selected),
                        assignments,
                        sample_size=min(1000, len(selected)),
                        random_state=42,
                    )
                )
            except ValueError:
                warnings.append("Silhouette score is unavailable for this cluster distribution.")
        metrics = {
            "cluster_sizes": {str(index): int(size) for index, size in enumerate(sizes)},
            "silhouette_score": silhouette,
            "inertia": float(model.inertia_),
        }
        predictions = [
            {"row": index, "cluster": int(cluster)}
            for index, cluster in enumerate(assignments[:100])
        ]
        return {
            "pipeline": pipeline,
            "metrics": metrics,
            "warnings": warnings,
            "rows": len(selected),
            "train_rows": len(selected),
            "test_rows": 0,
            "target": None,
            "task": task,
            "algorithm": algorithm,
            "features": feature_names,
            "parameters": parameters or {},
            "chart": predictions,
        }

    if task == "regression":
        if not pd.api.types.is_numeric_dtype(frame[target]):
            raise MLLabError("Regression requires a numeric target column.")
        target_values = pd.to_numeric(frame[target], errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )
    else:
        target_values = frame[target]
        if target_values.nunique(dropna=True) > 50:
            raise MLLabError("Classification target has more than 50 classes; choose a categorical target.")
    usable_target = target_values.notna()
    selected = selected.loc[usable_target].reset_index(drop=True)
    target_values = target_values.loc[usable_target].reset_index(drop=True)
    if len(selected) < 20:
        raise MLLabError("At least 20 complete target rows are required for a train/test evaluation.")
    if task == "classification":
        class_counts = target_values.value_counts()
        class_count = len(class_counts)
        test_count = max(class_count, int(math.ceil(len(target_values) * test_size)))
        if class_count < 2 or int(class_counts.min()) < 2 or len(target_values) - test_count < class_count:
            raise MLLabError("Each classification class needs at least two rows and enough samples for both splits.")
        split_kwargs = {"stratify": target_values}
        labels = sorted(target_values.unique().tolist(), key=str)
    else:
        test_count = max(1, int(math.ceil(len(target_values) * test_size)))
        split_kwargs = {}
        labels = None
    date_columns = [name for name in feature_names if name.lower() == "date" or name.lower().endswith("_date")]
    chronological = dataset_name == "Synthetic_India_Prices" and bool(date_columns)
    split_date_range = None
    if chronological:
        date_values = selected[date_columns[0]]
        unique_dates = np.sort(date_values.dropna().unique())
        test_date_count = max(1, int(math.ceil(len(unique_dates) * test_size)))
        if len(unique_dates) <= test_date_count:
            raise MLLabError("There are not enough distinct dates for a chronological train/test split.")
        cutoff = unique_dates[len(unique_dates) - test_date_count]
        train_mask = date_values < cutoff
        test_mask = date_values >= cutoff
        x_train, x_test = selected.loc[train_mask], selected.loc[test_mask]
        y_train, y_test = target_values.loc[train_mask], target_values.loc[test_mask]
        split_date_range = {
            "train": [
                pd.to_datetime(x_train[date_columns[0]].min(), unit="D").date().isoformat(),
                pd.to_datetime(x_train[date_columns[0]].max(), unit="D").date().isoformat(),
            ],
            "test": [
                pd.to_datetime(x_test[date_columns[0]].min(), unit="D").date().isoformat(),
                pd.to_datetime(x_test[date_columns[0]].max(), unit="D").date().isoformat(),
            ],
        }
        warnings.append("Chronological date split used: training rows precede the held-out test dates.")
    else:
        x_train, x_test, y_train, y_test = train_test_split(
            selected,
            target_values,
            test_size=test_count,
            random_state=42,
            **split_kwargs,
        )
    if len(x_train) < 2 or len(x_test) < 1:
        raise MLLabError("The selected train/test split is too small to evaluate.")
    pipeline = Pipeline([("preprocessor", _preprocessor(x_train)), ("model", model)])
    pipeline.fit(x_train, y_train)
    predicted = pipeline.predict(x_test)
    if task == "regression":
        mse = float(mean_squared_error(y_test, predicted))
        metrics = {
            "mae": float(mean_absolute_error(y_test, predicted)),
            "mse": mse,
            "rmse": float(math.sqrt(mse)),
            "r2": float(r2_score(y_test, predicted)) if len(y_test) >= 2 else None,
        }
        if len(y_test) < 10:
            warnings.append("The held-out test set is small; regression metrics may be unstable.")
    else:
        matrix = confusion_matrix(y_test, predicted, labels=labels)
        metrics = {
            "accuracy": float(accuracy_score(y_test, predicted)),
            "precision_weighted": float(precision_score(y_test, predicted, average="weighted", zero_division=0)),
            "recall_weighted": float(recall_score(y_test, predicted, average="weighted", zero_division=0)),
            "f1_weighted": float(f1_score(y_test, predicted, average="weighted", zero_division=0)),
            "labels": [str(label) for label in labels],
            "confusion_matrix": matrix.astype(int).tolist(),
        }
    chart = [
        {
            "actual": _json_value(actual),
            "predicted": _json_value(prediction),
        }
        for actual, prediction in zip(y_test.iloc[:100], predicted[:100])
    ]
    return {
        "pipeline": pipeline,
        "metrics": metrics,
        "warnings": warnings,
        "rows": len(selected),
        "train_rows": len(x_train),
        "test_rows": len(x_test),
        "target": target,
        "task": task,
        "algorithm": algorithm,
        "features": feature_names,
        "parameters": parameters or {},
        "chronological": chronological,
        "split_date_range": split_date_range,
        "chart": chart,
    }


def save_artifact(pipeline, artifact_id):
    if not re.fullmatch(r"[a-f0-9]{32}", artifact_id):
        raise MLLabError("Invalid model artifact identifier.")
    artifact_dir = Path(ARTIFACT_DIR)
    artifact_dir.mkdir(parents=True, exist_ok=True)
    destination = artifact_dir / f"{artifact_id}.joblib"
    temporary = artifact_dir / f".{artifact_id}.{uuid.uuid4().hex}.tmp"
    joblib.dump(pipeline, temporary, compress=3)
    os.replace(temporary, destination)
    return destination.name


def load_artifact(artifact_id):
    if not isinstance(artifact_id, str) or not re.fullmatch(r"[a-f0-9]{32}", artifact_id):
        raise MLLabError("The saved model is unavailable.")
    artifact_dir = Path(ARTIFACT_DIR).resolve()
    path = (artifact_dir / f"{artifact_id}.joblib").resolve()
    if path.parent != artifact_dir or not path.is_file():
        raise MLLabError("The saved model artifact is missing.")
    return joblib.load(path)


def prediction_features(dataset_name, feature_names):
    if dataset_name not in DATASET_META:
        raise MLLabError("This model's source dataset is unavailable.")
    frame = load_workbook()[dataset_name]
    features = []
    for name in feature_names:
        if name not in frame.columns:
            raise MLLabError("A model feature no longer exists in its source dataset.")
        values = frame[name]
        is_date = name.lower() == "date" or name.lower().endswith("_date")
        numeric = pd.api.types.is_numeric_dtype(values)
        non_missing = values.dropna()
        features.append(
            {
                "name": name,
                "kind": "date" if is_date else "numeric" if numeric else "categorical",
                "required": True,
                "minimum": float(non_missing.min()) if numeric and len(non_missing) else None,
                "maximum": float(non_missing.max()) if numeric and len(non_missing) else None,
                "options": [] if numeric or is_date else [str(value) for value in non_missing.astype(str).unique()[:100]],
            }
        )
    return features


def predict(pipeline, feature_specs, values):
    if not isinstance(values, dict) or set(values) != {item["name"] for item in feature_specs}:
        raise MLLabError("Supply exactly one input value for each model feature.")
    row = {}
    for feature in feature_specs:
        name = feature["name"]
        value = values[name]
        if feature["kind"] == "date":
            parsed_date = pd.to_datetime(value, errors="coerce", format="mixed")
            if pd.isna(parsed_date):
                raise MLLabError(f"{name} must be a valid date.")
            row[name] = parsed_date.strftime("%Y-%m-%d")
        elif feature["kind"] == "numeric":
            try:
                parsed = float(value)
            except (TypeError, ValueError):
                raise MLLabError(f"{name} must be a number.")
            if not math.isfinite(parsed):
                raise MLLabError(f"{name} must be a finite number.")
            row[name] = parsed
        else:
            parsed = str(value).strip()
            if len(parsed) > 200:
                raise MLLabError(f"{name} is too long.")
            row[name] = parsed if parsed else np.nan
    frame = _prepare_dates(pd.DataFrame([row]))
    try:
        predicted = pipeline.predict(frame)[0]
    except (ValueError, TypeError) as error:
        raise MLLabError(f"Prediction could not be calculated: {error}") from error
    result = {"prediction": _json_value(predicted)}
    model = pipeline.named_steps["model"]
    if hasattr(model, "predict_proba"):
        probabilities = pipeline.predict_proba(frame)[0]
        classes = getattr(model, "classes_", range(len(probabilities)))
        result["probabilities"] = [
            {"class": str(label), "probability": float(probability)}
            for label, probability in zip(classes, probabilities)
        ]
    if hasattr(model, "cluster_centers_"):
        result["note"] = "Cluster numbers are labels assigned by this model, not inherently meaningful categories."
    return result


def encode_json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, default=_json_value)
